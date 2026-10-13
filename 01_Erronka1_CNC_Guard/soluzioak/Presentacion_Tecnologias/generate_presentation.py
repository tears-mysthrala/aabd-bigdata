"""Export the source-of-truth CNC Guard slides to editable PPTX + matched PDF.

Only presentation assets are touched. Training, services and deployments are not run.
"""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

HERE = Path(__file__).resolve().parent
WIDTH, HEIGHT = 960, 540
NAVY, PANEL, WHITE, MUTE, TEAL = "0A1424", "162740", "F5F8FC", "B5C8DB", "55E3C5"
ACCENTS = ["55E3C5", "59BDF2", "EABD66"]
FONT = "DejaVu Sans"


def inches(n):
    return Inches(n / 72)


def add_box(slide, x, y, w, h, color, *, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        inches(x), inches(y), inches(w), inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()


def add_text(slide, x, y, w, h, value, size, color, bold=False, align=None):
    shape = slide.shapes.add_textbox(inches(x), inches(y), inches(w), inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    para = tf.paragraphs[0]
    para.text = value
    para.font.name = FONT
    para.font.size = Pt(size)
    para.font.bold = bold
    para.font.color.rgb = RGBColor.from_string(color)
    para.line_spacing = 1.2
    if align:
        para.alignment = align


def pdf_rect(c, x, y, w, h, color, radius=0):
    c.setFillColor(HexColor("#" + color))
    if radius:
        c.roundRect(x, HEIGHT-y-h, w, h, radius, fill=1, stroke=0)
    else:
        c.rect(x, HEIGHT-y-h, w, h, fill=1, stroke=0)


def pdf_text(c, x, y, value, size, color, bold=False):
    c.setFont("DJS-Bold" if bold else "DJS", size)
    c.setFillColor(HexColor("#" + color))
    c.drawString(x, HEIGHT-y-size, value)


def setup_fonts():
    folder = Path("/usr/share/fonts/truetype/dejavu")
    pdfmetrics.registerFont(TTFont("DJS", str(folder / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DJS-Bold", str(folder / "DejaVuSans-Bold.ttf")))


def fit_lines(value, width, height, start):
    size = float(start)
    while size > 8.5:
        lines = simpleSplit(value, "DJS", size, width)
        if len(lines) * size * 1.22 < height:
            return size, lines
        size -= 0.5
    return 8.5, simpleSplit(value, "DJS", 8.5, width)


def main():
    setup_fonts()
    info = json.loads((HERE / "contenido.json").read_text(encoding="utf-8"))
    slides = info["slides"]
    assert len(slides) == 33, "Refuse to export an incomplete slide deck"
    ppt = Presentation()
    ppt.slide_width, ppt.slide_height = Inches(13.333333), Inches(7.5)
    pdf = canvas.Canvas(str(HERE / "CNC_Guard_Tecnologias.pdf"), pagesize=(WIDTH, HEIGHT))
    pdf.setTitle("CNC Guard - Propuesta tecnologica 2026")
    for index, data in enumerate(slides, 1):
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        add_box(slide, 0, 0, WIDTH, HEIGHT, NAVY)
        pdf_rect(pdf, 0, 0, WIDTH, HEIGHT, NAVY)
        add_box(slide, 0, 0, 7, HEIGHT, TEAL)
        pdf_rect(pdf, 0, 0, 7, HEIGHT, TEAL)
        hdr = "CNC GUARD / PROPUESTA TECNOLÓGICA"
        add_text(slide, 50, 20, 520, 20, hdr, 8.5, TEAL, True)
        pdf_text(pdf, 50, 20, hdr, 8.5, TEAL, True)
        title = data["title"]
        tsize = 33 if len(title) < 43 else 28
        add_text(slide, 50, 55, 860, 60, title, tsize, WHITE, True)
        pdf_text(pdf, 50, 55, title, tsize, WHITE, True)
        lead = data.get("lead", "")
        add_text(slide, 50, 114, 855, 40, lead, 14, MUTE)
        for k, line in enumerate(simpleSplit(lead, "DJS", 14, 855)[:2]):
            pdf_text(pdf, 50, 114 + k * 18, line, 14, MUTE)
        add_box(slide, 50, 158, 860, 2.5, TEAL)
        pdf_rect(pdf, 50, 158, 860, 2.5, TEAL)
        rows = data.get("rows", [])
        n = len(rows)
        cols = n if n <= 3 else 2
        nr = 1 if n <= 3 else (n + 1) // 2
        top = 211 if index == 1 else 176
        gap, bottom = 16, 65
        pw = (860 - gap * (cols - 1)) / cols if n else 860
        ph = (HEIGHT - top - bottom - gap * (nr - 1)) / nr if n else 290
        for j, (label, value) in enumerate(rows):
            x = 50 + j % cols * (pw + gap)
            y = top + j // cols * (ph + gap)
            accent = ACCENTS[j % len(ACCENTS)]
            pdf_rect(pdf, x, y, pw, ph, PANEL, 9)
            add_box(slide, x, y, pw, ph, PANEL, radius=True)
            pdf_rect(pdf, x, y, 4, ph, accent)
            add_box(slide, x, y, 4, ph, accent)
            pad = 16 if n <= 3 else 12
            inner = pw - 2 * pad
            label_size = 14.5 if n <= 3 else 11.5
            pdf_text(pdf, x+pad, y+11, label, label_size, accent, True)
            add_text(slide, x+pad, y+10, inner, 30, label, label_size, accent, True)
            bodytop = 47 if n <= 3 else 29
            size, wrapped = fit_lines(value, inner, ph - bodytop - pad, 18.0 if n <= 3 else 12.5)
            add_text(slide, x+pad, y+bodytop, inner, ph-bodytop-pad, value, size, WHITE)
            for k, line in enumerate(wrapped):
                pdf_text(pdf, x+pad, y+bodytop+k*size*1.22, line, size, WHITE)
        req = (data.get("req") or "DISEÑO Y ALCANCE").strip()[:125]
        pdf_rect(pdf, 50, 487, 860, 1.0, "29425D")
        add_box(slide, 50, 487, 860, 1.0, "29425D")
        pdf_text(pdf, 51, 497, req, 8.0, MUTE)
        add_text(slide, 51, 497, 775, 20, req, 8.0, MUTE)
        pdf_text(pdf, 862, 497, f"{index:02d} / 33", 8.0, TEAL, True)
        add_text(slide, 844, 497, 65, 20, f"{index:02d} / 33", 8.0, TEAL,
                 True, PP_ALIGN.RIGHT)
        try:
            slide.notes_slide.notes_text_frame.text = data.get("note", "")
        except (AttributeError, ValueError):
            pass
        pdf.showPage()
    ppt.save(str(HERE / "CNC_Guard_Tecnologias.pptx"))
    pdf.save()
    print(f"Export OK: {len(slides)} editable PPTX slides + {len(slides)} PDF pages")


if __name__ == "__main__":
    main()
