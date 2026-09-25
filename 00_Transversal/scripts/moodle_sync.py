#!/usr/bin/env python3
"""Moodle fitxategien sinkronizazio automatikoa.

Orduro exekutatzen da (systemd timer bidez):
1. Moodle-ra konektatzen da eta saioa hasten du.
2. Atal guztiak aztertzen ditu (fitxategiak, karpetak, zereginak).
3. Aldaketak edo fitxategi berriak badaude, deskargatu eta dagokion karpetan jartzen ditu.
4. .docx fitxategiak .md formatura bihurtzen ditu.
5. Git commit eta push egiten ditu GitHub-era.
6. Mahaigaineko jakinarazpena bidaltzen du (notify-send).
"""

import os
import re
import sys
import time
import subprocess
import urllib.parse
from datetime import datetime
from pathlib import Path
import requests
from bs4 import BeautifulSoup

REPO_ROOT = Path("/home/tears/bigdata")
LOG_FILE = REPO_ROOT / "moodle_sync.log"

USERNAME = os.environ.get("MOODLE_USER", "")
PASSWORD = os.environ.get("MOODLE_PASS", "")
BASE_URL = "https://elearning20.hezkuntza.net/012053"
COURSE_ID = "535"


def log(msg: str) -> None:
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    entry = f"[{now}] {msg}"
    print(entry)
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(entry + "\n")


def notify(title: str, msg: str) -> None:
    try:
        subprocess.run(["notify-send", "-a", "Moodle Sync", title, msg], check=False)
    except Exception as e:
        log(f"Ezin izan da notify-send exekutatu: {e}")


def lortu_saioa() -> requests.Session:
    """Chromium bidez saioa hasi eta cookies-ak eskuratu."""
    if not USERNAME or not PASSWORD:
        raise SystemExit("Falta MOODLE_USER/MOODLE_PASS en el entorno (nunca en código).")
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch(
            executable_path="/usr/bin/chromium",
            headless=True,
            args=["--no-sandbox", "--disable-dev-shm-usage"],
        )
        page = browser.new_page()
        page.goto(f"{BASE_URL}/login/index.php")
        page.fill("#username", USERNAME)
        page.fill("#password", PASSWORD)
        page.click("#loginbtn")
        page.wait_for_load_state("networkidle")
        cookies = page.context.cookies()
        browser.close()

    session = requests.Session()
    for c in cookies:
        session.cookies.set(c["name"], c["value"], domain=c["domain"], path=c["path"])

    # Verificar autenticación real: sin sesión válida Moodle devuelve el
    # formulario de login (HTTP 200) y el escaneo vería 0 actividades,
    # informando "sin novedades" en falso. Fallar ruidoso en ese caso.
    probe = session.get(f"{BASE_URL}/course/view.php?id={COURSE_ID}", timeout=15)
    if 'id="username"' in probe.text or "login/index.php" in probe.url:
        raise RuntimeError("Login Moodle fallido: sesión no autenticada (¿contraseña cambiada?).")
    return session


def docx_to_md(docx_path: Path, md_path: Path) -> None:
    """Bihurtu docx testu aberatsa markdown garbi batera."""
    try:
        import docx

        doc = docx.Document(docx_path)
        lines = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if not txt:
                continue
            if p.style.name.startswith("Heading 1"):
                lines.append(f"# {txt}\n")
            elif p.style.name.startswith("Heading 2"):
                lines.append(f"## {txt}\n")
            elif p.style.name.startswith("Heading 3"):
                lines.append(f"### {txt}\n")
            else:
                lines.append(f"{txt}\n")
        for t in doc.tables:
            for row in t.rows:
                row_txt = [c.text.strip().replace("\n", " ") for c in row.cells]
                lines.append("| " + " | ".join(row_txt) + " |\n")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        log(f"Bihurtuta: {docx_path.name} -> {md_path.name}")
    except Exception as e:
        log(f"Errorea docx bihurtzean ({docx_path}): {e}")


def helburu_direktorioa(sec_id: int, izena: str, karpeta_izena: str = "") -> Path:
    """Zehaztu fitxategiaren helburu-karpeta atal eta fitxategi-izenaren arabera."""
    izena_lower = izena.lower()
    
    # 1. Karpeta espezifikoak
    if "mock datuak" in karpeta_izena.lower():
        return REPO_ROOT / "04_Programazioa_5073" / "data" / "mock_datuak"
    if "alborapenak" in karpeta_izena.lower():
        return REPO_ROOT / "02_AA_Ereduak_5071" / "materialak" / "Alborapenak"

    # 2. Atalen araberako sailkapena
    if sec_id in (0, 1):
        return REPO_ROOT / "00_Transversal" / "materialak_00_Orokorra"
    elif sec_id in (2, 3):
        return REPO_ROOT / "01_Erronka1_CNC_Guard" / "materialak"
    elif sec_id == 4:
        return REPO_ROOT / "02_AA_Ereduak_5071" / "materialak"
    elif sec_id == 5:
        return REPO_ROOT / "03_ML_5072" / "materialak"
    elif sec_id == 6:
        return REPO_ROOT / "04_Programazioa_5073" / "materialak"
    elif sec_id == 7:
        return REPO_ROOT / "05_BigData_Ingeniaritza" / "materialak"
    elif sec_id == 8:
        if "kafka" in izena_lower or "01_03" in izena_lower:
            return REPO_ROOT / "07_Kafka" / "materialak"
        return REPO_ROOT / "06_NiFi" / "materialak"
    elif sec_id == 9:
        return REPO_ROOT / "08_Erronka2_RetailEUS" / "materialak"
    elif sec_id == 10:
        return REPO_ROOT / "00_Transversal" / "azterketak"
    else:
        return REPO_ROOT / "materialak"


def sinkronizatu() -> list[str]:
    """Exekutatu Moodle-ko eskaneatze eta deskarga prozesu osoa."""
    session = lortu_saioa()
    deskargatutakoak: list[str] = []

    log("Moodle atalak aztertzen...")
    for sec_id in range(0, 13):
        url = f"{BASE_URL}/course/view.php?id={COURSE_ID}&section={sec_id}"
        resp = session.get(url, timeout=15)
        if resp.status_code != 200:
            continue

        soup = BeautifulSoup(resp.text, "html.parser")
        activities = soup.select(".activity")

        for act in activities:
            link = act.select_one("a")
            if not link or not link.get("href"):
                continue

            act_url = link["href"]
            inst_el = act.select_one(".instancename") or link
            act_izena = inst_el.text.strip().split("\n")[0].strip()

            # 1. Baliabide zuzenak (mod/resource)
            if "mod/resource" in act_url:
                try:
                    head = session.head(act_url, allow_redirects=True, timeout=10)
                    cd = head.headers.get("Content-Disposition", "")
                    fname = None
                    if "filename=" in cd:
                        fname = cd.split("filename=")[1].strip("\"' ")
                    elif "/" in head.url:
                        fname = head.url.split("/")[-1].split("?")[0]

                    if not fname:
                        continue

                    fname = urllib.parse.unquote(fname)
                    remote_size = int(head.headers.get("Content-Length") or 0)
                    dest_dir = helburu_direktorioa(sec_id, fname)
                    dest_path = dest_dir / fname

                    behar_da = False
                    if not dest_path.exists():
                        behar_da = True
                    elif remote_size > 0 and dest_path.stat().st_size != remote_size:
                        behar_da = True

                    if behar_da:
                        dest_dir.mkdir(parents=True, exist_ok=True)
                        log(f"Deskargatzen: [Sec {sec_id}] {fname} ({remote_size} bytes)...")
                        r_file = session.get(act_url, allow_redirects=True)
                        with open(dest_path, "wb") as f:
                            f.write(r_file.content)
                        deskargatutakoak.append(str(dest_path.relative_to(REPO_ROOT)))

                        if fname.lower().endswith(".docx"):
                            md_path = dest_path.with_suffix(".md")
                            docx_to_md(dest_path, md_path)
                            deskargatutakoak.append(str(md_path.relative_to(REPO_ROOT)))

                except Exception as e:
                    log(f"Errorea baliabidea aztertzean ({act_url}): {e}")

            # 2. Karpetak (mod/folder)
            elif "mod/folder" in act_url:
                try:
                    r_fld = session.get(act_url, timeout=15)
                    f_soup = BeautifulSoup(r_fld.text, "html.parser")
                    base_dest = helburu_direktorioa(sec_id, "", karpeta_izena=act_izena)

                    for a in f_soup.find_all("a"):
                        fhref = a.get("href", "")
                        if "pluginfile.php" in fhref:
                            # Estrakzio bidea URL-tik (adib. Ariketa%203.1/ikasleak_notak_100.csv)
                            match = re.search(r"content/\d+/(.+?)\?", fhref)
                            erlatiboa = match.group(1) if match else a.text.strip()
                            erlatiboa = urllib.parse.unquote(erlatiboa)
                            dest_path = base_dest / erlatiboa

                            head = session.head(fhref, allow_redirects=True, timeout=10)
                            remote_size = int(head.headers.get("Content-Length") or 0)

                            behar_da = False
                            if not dest_path.exists():
                                behar_da = True
                            elif remote_size > 0 and dest_path.stat().st_size != remote_size:
                                behar_da = True

                            if behar_da:
                                dest_path.parent.mkdir(parents=True, exist_ok=True)
                                log(f"Deskargatzen karpetatik: {erlatiboa} ({remote_size} bytes)...")
                                r_file = session.get(fhref, allow_redirects=True)
                                with open(dest_path, "wb") as f:
                                    f.write(r_file.content)
                                deskargatutakoak.append(str(dest_path.relative_to(REPO_ROOT)))

                                if dest_path.name.lower().endswith(".docx"):
                                    md_path = dest_path.with_suffix(".md")
                                    docx_to_md(dest_path, md_path)
                                    deskargatutakoak.append(str(md_path.relative_to(REPO_ROOT)))

                except Exception as e:
                    log(f"Errorea karpeta aztertzean ({act_url}): {e}")

            # 3. Zereginen eranskinak (mod/assign)
            elif "mod/assign" in act_url:
                try:
                    r_asg = session.get(act_url, timeout=15)
                    asg_soup = BeautifulSoup(r_asg.text, "html.parser")
                    for a in asg_soup.select("#intro a, .introattachment a, .fileuploadsubmission a"):
                        ahref = a.get("href", "")
                        if "pluginfile.php" in ahref and "forcedownload=1" in ahref:
                            fname = a.text.strip()
                            if not fname:
                                continue
                            fname = urllib.parse.unquote(fname)
                            dest_dir = helburu_direktorioa(sec_id, fname)
                            dest_path = dest_dir / fname

                            head = session.head(ahref, allow_redirects=True, timeout=10)
                            remote_size = int(head.headers.get("Content-Length") or 0)

                            behar_da = False
                            if not dest_path.exists():
                                behar_da = True
                            elif remote_size > 0 and dest_path.stat().st_size != remote_size:
                                behar_da = True

                            if behar_da:
                                dest_dir.mkdir(parents=True, exist_ok=True)
                                log(f"Deskargatzen zereginetik: {fname} ({remote_size} bytes)...")
                                r_file = session.get(ahref, allow_redirects=True)
                                with open(dest_path, "wb") as f:
                                    f.write(r_file.content)
                                deskargatutakoak.append(str(dest_path.relative_to(REPO_ROOT)))

                                if fname.lower().endswith(".docx"):
                                    md_path = dest_path.with_suffix(".md")
                                    docx_to_md(dest_path, md_path)
                                    deskargatutakoak.append(str(md_path.relative_to(REPO_ROOT)))
                except Exception as e:
                    log(f"Errorea zeregina aztertzean ({act_url}): {e}")

    return deskargatutakoak


SAIAKERA_MAX = 3
SAIAKERA_ATSEDENA_S = 30


def main() -> None:
    log("=== Sinkronizazio zikloa hasita ===")
    berriak: list[str] = []
    try:
        for saiakera in range(1, SAIAKERA_MAX + 1):
            try:
                berriak = sinkronizatu()
                break
            except Exception as e:
                log(f"{saiakera}. saiakera huts: {e}")
                if saiakera >= SAIAKERA_MAX:
                    raise
                time.sleep(SAIAKERA_ATSEDENA_S)
        if berriak:
            log(f"Deskargatutako fitxategiak ({len(berriak)}): {', '.join(berriak)}")
            notify(
                "Moodle Eguneratua!",
                f"{len(berriak)} fitxategi berri deskargatu dira:\n" + "\n".join(berriak[:3]),
            )

            # Git: bidea ematen dugu deskargatutako fitxategiei soilik.
            # Ez erabili `git add -A`: lan-arloko aldaketa ajenoak ez dira
            # inoiz "Auto-sync" commit batean sartu behar.
            subprocess.run(["git", "add", "--", *berriak], cwd=REPO_ROOT, check=True)
            staged = subprocess.run(
                ["git", "diff", "--cached", "--quiet"],
                cwd=REPO_ROOT,
            )
            if staged.returncode != 0:
                msg = f"Auto-sync Moodle: {len(berriak)} fitxategi deskargatuta\n\n" + "\n".join(f"- {b}" for b in berriak)
                subprocess.run(["git", "commit", "-m", msg], cwd=REPO_ROOT, check=True)
                subprocess.run(["git", "push", "origin", "master"], cwd=REPO_ROOT, check=True)
                log("Git push ondo osatu da.")
            else:
                log("Aldaketarik ez stage-an: commit/push ez da egin.")
        else:
            log("Ez dago fitxategi berririk Moodle-n.")

    except Exception as e:
        log(f"Errore orokorra sinkronizazioan: {e}")
        notify("Moodle Sync Errorea", str(e))
        sys.exit(1)


if __name__ == "__main__":
    main()
